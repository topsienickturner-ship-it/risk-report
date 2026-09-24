#!/usr/bin/env python3
"""Generate a landscape risk report from an Excel risk/action export."""
import argparse
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import sys
import re
import tempfile
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate, Frame, LongTable, PageBreak, PageTemplate, Paragraph,
    Spacer, Table, TableStyle,
)


@dataclass
class Risk:
    risk_id: str
    details: dict = field(default_factory=dict)
    actions: list = field(default_factory=list)


def text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    return str(value).strip()


# Match the export's explicit two-digit marker, not numeric prose or IDs.
ACTION_PREFIX = re.compile(r"(?m)^[^\S\n]*([0-9]{2})[^\S\n]+-(?:[^\S\n]+|$)")


def split_actions(record):
    """Align numbered entries by their prefix, not by their position in a cell."""
    fields = {k: v for k, v in record.items() if k.startswith('Action ')}
    parsed = {}
    numbers = set()
    for key, value in fields.items():
        value = '\n'.join(value.splitlines())
        matches = list(ACTION_PREFIX.finditer(value))
        if not matches:
            continue
        if value[:matches[0].start()].strip():
            raise ValueError(f'{key}: text before first numbered action.')
        entries = {}
        for i, match in enumerate(matches):
            number = int(match.group(1))
            if number in entries:
                raise ValueError(f'{key}: repeated action prefix {number:02d}.')
            end = matches[i+1].start() if i+1 < len(matches) else len(value)
            entries[number] = value[match.end():end].strip()
        parsed[key] = entries
        numbers.update(entries)
    if not numbers:
        return [fields] if any(fields.values()) else []
    if len(numbers) > 1:
        for key, value in fields.items():
            if value and key not in parsed:
                raise ValueError(f'{key}: unnumbered value alongside multiple numbered actions; '
                                 'prefix each value to identify its action.')
    actions = []
    for number in sorted(numbers):
        action = {key: parsed[key].get(number, '') if key in parsed else value
                  for key, value in fields.items()}
        action['_number'] = f'{number:02d}'
        actions.append(action)
    return actions


def read_risks(path, sheet=None, header_row=None):
    """Keep first nonempty shared values; reject ambiguous or unassigned data."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[sheet] if sheet else workbook.active
        # Some exporters write incorrect worksheet dimensions. Read the actual
        # cells and retain column positions, including blank leading columns.
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        headers = []
        for source_row, values in enumerate(rows, 1):
            candidate = [text(v) for v in values]
            if (header_row is None and 'Risk ID' in candidate) or source_row == header_row:
                headers = candidate
                header_row = source_row
                break
        if 'Risk ID' not in headers:
            raise ValueError(f"No 'Risk ID' column found in {ws.title!r}"
                             + (f' on row {header_row}.' if header_row else '.'))
        named = [h for h in headers if h]
        if len(named) != len(set(named)):
            raise ValueError('Duplicate column headings; each named column must be unique.')
        risks = {}
        seen_actions = {}
        for row_number, values in enumerate(rows, header_row + 1):
            record = {h: text(v) for h, v in zip(headers, values) if h}
            if not any(record.values()):
                continue
            risk_id = record.get('Risk ID', '')
            if not risk_id:
                raise ValueError(f'Row {row_number}: populated row has no Risk ID.')
            risk = risks.setdefault(risk_id, Risk(risk_id))
            for key, value in record.items():
                if key.startswith('Action ') or not value:
                    continue
                previous = risk.details.get(key)
                if previous and previous != value:
                    raise ValueError(f'Row {row_number}, risk {risk_id}: conflicting {key!r}: '
                                     f'{previous!r} / {value!r}.')
                risk.details[key] = value
            try:
                actions = split_actions(record)
            except ValueError as exc:
                raise ValueError(f'Row {row_number}, risk {risk_id}: {exc}') from exc
            for action in actions:
                action_id = action.get('Action ID')
                if action_id:
                    key = (risk_id, action_id)
                    comparable = {k: v for k, v in action.items() if k != '_number'}
                    if key in seen_actions:
                        if seen_actions[key] != comparable:
                            raise ValueError(f'Row {row_number}: conflicting duplicate action '
                                             f'{action_id!r} for risk {risk_id}.')
                        continue
                    seen_actions[key] = comparable
                risk.actions.append(action)
        if not risks:
            raise ValueError('No risk records found beneath the column headings.')
        return list(risks.values())
    finally:
        workbook.close()


BODY = ParagraphStyle('Body', fontName='Helvetica', fontSize=9, leading=12,
                      spaceAfter=5)
SMALL = ParagraphStyle('Small', parent=BODY, fontSize=8, leading=10, spaceAfter=0)
HEADING = ParagraphStyle('Heading', parent=BODY, fontName='Helvetica-Bold',
                         fontSize=13, leading=16, spaceAfter=10, keepWithNext=True)


def para(value, style=BODY):
    return Paragraph(escape(text(value) or '—').replace('\n', '<br/>'), style)


def labelled(label, value):
    return Paragraph(f'<b>{escape(label)}:</b> '
                     + escape(text(value) or '—').replace('\n', '<br/>'), BODY)


class RiskHeading(Paragraph):
    def __init__(self, risk):
        super().__init__(escape(f'{risk.risk_id} — {risk.details.get("Title", "")}'), HEADING)
        self.risk = risk


class RiskDocument(BaseDocTemplate):
    def __init__(self, filename, project, report_date):
        super().__init__(str(filename), pagesize=landscape(A4),
                         leftMargin=14*mm, rightMargin=14*mm,
                         topMargin=28*mm, bottomMargin=16*mm,
                         title=f'{project} Risk Register', author='Risk Report')
        self.project = project
        self.report_date = report_date
        self.current_risk = None
        self.risk_start_page = 0
        self.page_risks = []  # Useful for validating pagination without a PDF parser.
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height,
                      leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
        self.addPageTemplates(PageTemplate(id='risk', frames=[frame], onPageEnd=self.footer))

    def afterFlowable(self, flowable):
        if isinstance(flowable, RiskHeading):
            self.current_risk = flowable.risk
            self.risk_start_page = self.page

    def footer(self, canvas, doc):
        canvas.saveState()
        width, height = self.pagesize
        canvas.setFont('Helvetica-Bold', 14)
        canvas.drawCentredString(width/2, height-13*mm, self.project + ' RISK REGISTER')
        canvas.setFont('Helvetica', 8)
        canvas.drawRightString(width-self.rightMargin, height-20*mm, self.report_date)
        if self.current_risk:
            risk = self.current_risk
            suffix = ' (continued)' if self.page > self.risk_start_page else ''
            label = f'{risk.risk_id} — {risk.details.get("Title", "")}{suffix}'
            # A fixed-height header must not intrude into flowing report content.
            max_width = self.width - 110
            while canvas.stringWidth(label, 'Helvetica', 8) > max_width and len(label) > 4:
                label = label[:-4] + '...'
            canvas.drawString(self.leftMargin, height-20*mm, label)
            self.page_risks.append(risk.risk_id)
        canvas.setStrokeColor(colors.HexColor('#b7c1cc'))
        canvas.line(self.leftMargin, height-23*mm, width-self.rightMargin, height-23*mm)
        canvas.drawRightString(width-self.rightMargin, 9*mm, f'Page {self.page}')
        canvas.restoreState()


def table_style():
    return TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e7edf3')),
        ('LINEBELOW', (0, 0), (-1, 0), .6, colors.HexColor('#566778')),
        ('LINEBELOW', (0, 1), (-1, -1), .3, colors.HexColor('#ccd3db')),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ])


def build_report(risks, output, project='Project', report_date=None, cost='planned'):
    doc = RiskDocument(output, project, report_date or date.today().strftime('%d/%m/%Y'))
    story = []
    for index, risk in enumerate(risks):
        if index:
            story.append(PageBreak())
        d = risk.details
        story.append(RiskHeading(risk))
        for label, key in [('Status', 'Risk Status'), ('Owner', 'Owner'), ('Type', 'Type'),
                           ('Risk', 'Description'), ('Cause', 'Cause'), ('Effect', 'Impact')]:
            story.append(labelled(label, d.get(key)))
        story.append(Spacer(1, 3*mm))
        scores = [[para(x, SMALL) for x in
                   ['Assessment', 'Probability', 'Schedule', 'Cost', 'Performance', 'Score']]]
        for label, prefix, score_key in [('Initial', 'Init', 'Initial Score'),
                                         ('Current', 'Curr', 'Current Score'),
                                         ('Target', 'Target', 'Target Score')]:
            qual_prefix = 'Initial' if prefix == 'Init' else prefix
            performance = d.get(f'{prefix} Qual Tech/Eng Performance')
            if prefix == 'Target':
                performance = performance or d.get('Targetr Qual Tech/Eng Performance')
            scores.append([para(x, SMALL) for x in [label,
                d.get(f'{qual_prefix} Qual Prob'), d.get(f'{qual_prefix} Qual Sched'),
                d.get(f'{qual_prefix} Qual Cost'), performance, d.get(score_key)]])
        story.append(Table(scores, colWidths=[doc.width/6]*6, style=table_style()))
        story.append(Spacer(1, 5*mm))
        story.append(Paragraph('Actions', HEADING))
        if not risk.actions:
            story.append(para('No actions recorded.'))
            continue
        columns = [('Action ID', 'ID'), ('Action Title', 'Action'), ('Action Owner', 'Responsibility'),
                   ('Action Status', 'Status'), ('Action Start Date (forecast)', 'Start'),
                   ('Action Finish Date (forecast)', 'Finish'), ('Action last comment', 'Notes'),
                   ('Action Planned Spend' if cost == 'planned' else 'Action Actual Spend',
                    'Planned cost' if cost == 'planned' else 'Actual cost')]
        rows = [[para(label, SMALL) for _, label in columns]]
        rows.extend([[para((action.get(key) or action.get('_number'))
                               if key == 'Action ID' else action.get(key), SMALL)
                      for key, _ in columns] for action in risk.actions])
        story.append(LongTable(rows, colWidths=[doc.width*f for f in
                          [.06, .23, .13, .08, .08, .08, .25, .09]],
                          repeatRows=1, splitByRow=1, splitInRow=1, style=table_style()))
    doc.build(story)
    return doc.page_risks


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path, help='Excel .xlsx export')
    parser.add_argument('-o', '--output', type=Path, default=Path('risk-report.pdf'))
    parser.add_argument('--sheet', help='Worksheet name (default: active worksheet)')
    parser.add_argument('--header-row', type=int, help='Header row (default: detect Risk ID column)')
    parser.add_argument('--project', default='Project', help='Project title in page headers')
    parser.add_argument('--report-date', help='Printed date (default: today)')
    parser.add_argument('--cost', choices=['planned', 'actual'], default='planned')
    args = parser.parse_args()
    if args.header_row is not None and args.header_row < 1:
        parser.error('--header-row must be positive')
    if args.input.resolve() == args.output.resolve():
        parser.error('Input and output paths must differ')
    temporary = None
    try:
        risks = read_risks(args.input, args.sheet, args.header_row)
        # Render atomically so an error cannot leave a partially written report.
        with tempfile.NamedTemporaryFile(dir=args.output.parent, suffix='.pdf', delete=False) as f:
            temporary = Path(f.name)
        pages = build_report(risks, temporary, args.project, args.report_date, args.cost)
        temporary.replace(args.output)
    except (OSError, ValueError, KeyError) as exc:
        print(f'Error: {exc}', file=sys.stderr)
        return 1
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    print(f'Created {args.output}: {len(risks)} risks, {len(pages)} pages.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
